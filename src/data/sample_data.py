"""生成"Y企业风格"示例数据集（按 data/y_template.xlsx 表结构）。

【示例数据声明】本模块生成的数据为仿真示例数据，参数量级参考模板示例行
与 V-toy-1 合理外推，不代表 Y 企业真实数据；文件名与各表元信息均显式标注。

规模：6 产品族 / 8 产品型号 / 3 阶段 8 机器 / 10 计划周期 / 28 订单。
机制设计：τ=4 DIP 高负荷（加班型反馈）；τ=6 高技能工人缺勤（η=7/8 折减）。
"""
from __future__ import annotations

from datetime import date, timedelta

import openpyxl

START_DATE = date(2026, 3, 1)  # 排产起始日期；周期 = 起始日 + (t−1) 天

FAMILIES = ["F1", "F2", "F3", "F4", "F5", "F6"]
MODELS = {  # 型号 → (族, 型号名称)
    "KX-25A": ("F1", "25L电烤炉A"), "KX-25B": ("F1", "25L电烤炉B"),
    "KX-32A": ("F2", "32L电烤炉"), "QR-40A": ("F3", "40L蒸烤箱"),
    "DL-18A": ("F4", "18L小烤箱A"), "DL-18B": ("F4", "18L小烤箱B"),
    "QR-52A": ("F5", "52L蒸烤箱"), "PT-10A": ("F6", "10L迷你炉"),
}
# 族级参数：c_p, h_p, b_p, g_p, U_lot
COSTS = {
    "F1": (50, 2, 10, 100, 5), "F2": (60, 2, 9, 120, 5),
    "F3": (70, 3, 12, 150, 5), "F4": (55, 2, 8, 110, 6),
    "F5": (75, 3, 14, 160, 4), "F6": (45, 1, 8, 90, 6),
}
# 档位1 基准单件加工时间 (SMT, DIP, 总装)
PROC = {
    "F1": (8, 14, 10), "F2": (10, 18, 12), "F3": (7, 20, 11),
    "F4": (9, 16, 9), "F5": (12, 22, 14), "F6": (6, 12, 8),
}
THETA = {1: 1.00, 2: 0.90, 3: 0.80}
SDST = {  # 族级换模矩阵（min）；F3↔F5 为深度换模
    "F1": {"F1": 0, "F2": 25, "F3": 30, "F4": 20, "F5": 35, "F6": 15},
    "F2": {"F1": 25, "F2": 0, "F3": 30, "F4": 25, "F5": 40, "F6": 20},
    "F3": {"F1": 30, "F2": 30, "F3": 0, "F4": 25, "F5": 240, "F6": 25},
    "F4": {"F1": 20, "F2": 25, "F3": 25, "F4": 0, "F5": 30, "F6": 15},
    "F5": {"F1": 35, "F2": 40, "F3": 240, "F4": 30, "F5": 0, "F6": 30},
    "F6": {"F1": 15, "F2": 20, "F3": 25, "F4": 15, "F5": 30, "F6": 0},
}
INITIAL_SETUP = 10
STAGES = [(1, "SMT贴装", ["M1-1", "M1-2", "M1-3"]),
          (2, "DIP插件", ["M2-1", "M2-2"]),
          (3, "总装", ["M3-1", "M3-2", "M3-3"])]
# 订单流：(订单号, 型号, 数量, 到达周期, 交付周期)
ORDERS = [
    ("SO-001", "KX-25A", 12, 1, 1), ("SO-002", "KX-25B", 8, 1, 1),
    ("SO-003", "PT-10A", 15, 1, 1),
    ("SO-004", "KX-32A", 15, 1, 2), ("SO-005", "DL-18A", 18, 1, 2),
    ("SO-006", "KX-25A", 15, 1, 3), ("SO-007", "QR-40A", 12, 2, 3),
    ("SO-008", "PT-10A", 10, 1, 3),
    ("SO-009", "KX-32A", 20, 2, 4), ("SO-010", "DL-18B", 20, 3, 4),
    ("SO-011", "KX-25B", 12, 3, 4),
    ("SO-012", "QR-52A", 10, 3, 5), ("SO-013", "PT-10A", 20, 4, 5),
    ("SO-014", "KX-25A", 18, 4, 6), ("SO-015", "DL-18A", 15, 5, 6),
    ("SO-016", "QR-40A", 20, 5, 7), ("SO-017", "QR-52A", 16, 6, 7),
    ("SO-018", "KX-32A", 18, 6, 8), ("SO-019", "PT-10A", 18, 7, 8),
    ("SO-020", "QR-52A", 8, 7, 9), ("SO-021", "DL-18B", 12, 8, 9),
    ("SO-022", "KX-25B", 10, 8, 9),
    ("SO-023", "QR-40A", 10, 8, 10), ("SO-024", "PT-10A", 12, 9, 10),
    ("SO-025", "KX-25A", 6, 2, 5), ("SO-026", "DL-18A", 10, 4, 8),
    ("SO-027", "KX-32A", 8, 5, 9), ("SO-028", "PT-10A", 8, 3, 6),
]
WORKER_LEVELS = [(1, "初级技工", 1, 200, 30), (2, "中级技工", 2, 280, 38),
                 (3, "高级技工", 3, 400, 55)]
NORMAL_AVAIL = (4, 3, 2)
ABSENT_PERIOD = 6          # τ=6 高技能工人缺勤
ABSENT_AVAIL = (4, 3, 0)   # 总数 7 < 8 台机器 → η=7/8
T_MAX = 10

MARK = "【示例数据：仿真生成，非企业真实数据】"


def _d(t: int) -> str:
    return (START_DATE + timedelta(days=t - 1)).isoformat()


def build_sample_workbook(path: str) -> None:
    """生成示例数据工作簿（表结构与 y_template.xlsx 对齐）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "填写说明"
    ws.append(["Y企业风格示例数据集 " + MARK])
    ws.append(["用途：Stage 6 实验框架（方案A~D、消融、敏感性）的输入数据。"])
    ws.append(["规模：6 产品族 / 8 型号 / 3 阶段 8 机器 / 10 周期 / 28 订单。"])
    ws.append(["数据口径：结构沿用 y_template.xlsx；全部数值为仿真示例。"])

    ws = wb.create_sheet("一 订单与需求")
    ws.append(["一、订单与需求数据 " + MARK])
    ws.append(["表1-1 历史订单明细"])
    ws.append(["订单号", "产品型号", "数量（件）", "下单日期", "要求交付日期", "备注"])
    for oid, model, qty, arr, due in ORDERS:
        ws.append([oid, model, qty, _d(arr), _d(due), "示例订单"])
    ws.append([])
    ws.append(["表1-2 产品型号与产品族对照"])
    ws.append(["产品型号", "型号名称", "所属产品族", "备注（分族依据：外壳/工艺/换模关系）"])
    for model, (fam, name) in MODELS.items():
        ws.append([model, name, fam, "示例"])
    ws.append([])
    ws.append(["表1-3 期初库存与欠交状态"])
    ws.append(["产品型号", "期初库存（件）", "期初欠交（件）", "统计截止日期"])
    for model in MODELS:
        ws.append([model, 0, 0, _d(1)])

    ws = wb.create_sheet("二 成本参数")
    ws.append(["二、成本参数 " + MARK])
    ws.append(["表2-1 产品成本参数"])
    ws.append(["产品型号", "单位生产成本（元/件）", "单位库存持有成本（元/件·周期）",
               "单位延期惩罚成本（元/件·周期）", "生产启动成本（元/次）", "数据来源/口径说明"])
    for model, (fam, _) in MODELS.items():
        c, h, b, g, _u = COSTS[fam]
        ws.append([model, c, h, b, g, "示例数据"])

    ws = wb.create_sheet("三 车间结构与工艺")
    ws.append(["三、车间结构与工艺数据 " + MARK])
    ws.append(["表3-1 加工阶段与并行机配置"])
    ws.append(["阶段编号", "阶段名称", "机器编号", "机器型号", "可用速度档位数", "备注"])
    for j, name, machines in STAGES:
        for m in machines:
            ws.append([j, name, m, "示例机型", 3, "档位1低速/2标准/3高速"])
    ws.append([])
    ws.append(["表3-2 单位件加工时间（长表：一行一个组合）"])
    ws.append(["产品型号", "阶段编号", "机器编号", "速度档位", "单件加工时间（min/件）", "数据来源"])
    for model, (fam, _) in MODELS.items():
        for j, _, _ in STAGES:
            base = PROC[fam][j - 1]
            for s, theta in THETA.items():
                ws.append([model, j, "全部同型", s, round(base * theta, 2), "示例"])
    ws.append([])
    ws.append(["表3-3 序列依赖换模时间（按产品族、分阶段填报）"])
    ws.append(["阶段编号", "机器编号/机型", "前产品族", "后产品族", "换模时间（min）", "备注"])
    for fam in FAMILIES:
        ws.append(["全部", "全部同型", "开机", fam, INITIAL_SETUP, "初始设置"])
    for a in FAMILIES:
        for b in FAMILIES:
            if a != b:
                note = "深度换模（拆装模具）" if SDST[a][b] >= 200 else ""
                ws.append(["全部", "全部同型", a, b, SDST[a][b], note])
    ws.append([])
    ws.append(["表3-4 跨阶段运输时间"])
    ws.append(["前阶段机器", "后阶段机器", "运输时间（min）", "方式（传送带/人工/AGV）"])
    ws.append(["全部", "全部", 5, "传送带（全部取 5 min，示例）"])
    ws.append([])
    ws.append(["表3-6 子批批量上限"])
    ws.append(["产品型号", "最大子批批量（件）", "依据（栈板/周转箱/转运批量）"])
    for model, (fam, _) in MODELS.items():
        ws.append([model, COSTS[fam][4], "标准栈板容量（示例）"])

    ws = wb.create_sheet("四 能耗参数")
    ws.append(["四、能耗参数 " + MARK])
    ws.append(["表4-1 机器功率（分速度档位）"])
    ws.append(["阶段编号", "机器编号", "速度档位", "加工功率（kW）", "准备/换模功率（kW）", "空闲待机功率（kW）", "数据来源"])
    for s, pe in [(1, 3), (2, 5), (3, 8)]:
        ws.append(["全部", "全部同型", s, pe, 2, 1, "示例"])
    ws.append([])
    ws.append(["表4-2 其他能耗与电价"])
    ws.append(["项目", "数值", "单位", "备注"])
    ws.append(["工业电价", 0.85, "元/kWh", "示例（不进入调度模型）"])
    ws.append(["单次运输能耗", 0.01, "kWh/次", "示例"])
    ws.append(["车间辅助功率", 0.5, "kW", "示例"])

    ws = wb.create_sheet("五 人力资源")
    ws.append(["五、人力资源数据 " + MARK])
    ws.append(["表5-1 技能等级定义与工资"])
    ws.append(["技能等级", "等级名称", "可操作的最高速度档位", "班次工资（元/班）", "加班费率（元/小时）", "备注"])
    for lv, name, max_s, wage, ot_rate in WORKER_LEVELS:
        ws.append([lv, name, max_s, wage, ot_rate, "向下兼容"])
    ws.append([])
    ws.append(["表5-2 各周期可用工人数（按排班表）"])
    ws.append(["周期/日期", "等级1可用人数", "等级2可用人数", "等级3可用人数", "备注（请假/借调等）"])
    ws.append(["常态配置", *NORMAL_AVAIL, "示例"])
    ws.append([_d(ABSENT_PERIOD), *ABSENT_AVAIL, "高级技工全员缺勤（机制示例）"])

    ws = wb.create_sheet("六 日历与产能")
    ws.append(["六、日历与产能数据 " + MARK])
    ws.append(["表6-1 班次与工时配置"])
    ws.append(["项目", "数值", "单位", "备注"])
    for row in [
        ("每日班次数", 1, "班", "示例"),
        ("每班标准工时", 480, "min/班", "8小时制"),
        ("每周工作天数", 7, "天", "连续排产（示例简化）"),
        ("单日加班上限", 120, "min", "示例"),
        ("计划周期长度", 1, "天", "1 周期 = 1 天"),
        ("排产起始日期", START_DATE.isoformat(), "日期", "周期1 对应日"),
        ("计划周期数", T_MAX, "周期", "T_max"),
    ]:
        ws.append(list(row))
    wb.save(path)
