"""Excel 适配器：V-toy-1 金标工作簿读入器 + Y企业模板结构读入骨架。

V-toy-1（data/vtoy1.xlsx）为只读金标，版式已冻结（tests/test_vtoy1_workbook.py
按相同坐标锁定），故按单元格坐标解析；坐标常量与金标 sheet 一一对应。

Y企业模板（data/y_template.xlsx）本阶段只读结构（各 sheet 的"表x-y"标题与
表头），字段到 ProblemData 的映射留待 Stage 6 与示例数据一并对齐。
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .problem_data import EnergyParams, Order, ProblemData
from .validators import validate_problem_data

# ── 通用解析助手 ──────────────────────────────────────────────


def _parse_gear_map(text: str) -> dict[int, float]:
    """解析 's=1:1.00; s=2:0.90; s=3:0.80' 形式的档位映射。"""
    pairs = re.findall(r"s\s*=\s*(\d+)\s*[:：]\s*([\d.]+)", str(text))
    return {int(s): float(v) for s, v in pairs}


def _parse_int_set(text: str) -> list[int]:
    """解析 '{1,2,3}' 形式的整数集合（取第一个花括号内容）。"""
    inner = re.search(r"\{([\d,\s]+)\}", str(text)).group(1)
    return [int(x) for x in inner.split(",")]


def _parse_machines(text: str) -> dict[int, list[str]]:
    """解析 'M1={M11,M12}; M2={M21}; M3={M31,M32}' 为 {阶段: [机器]}。"""
    result: dict[int, list[str]] = {}
    for stage, inner in re.findall(r"M(\d+)\s*=\s*\{([^}]*)\}", str(text)):
        result[int(stage)] = [m.strip() for m in inner.split(",") if m.strip()]
    return result


def _parse_stage_names(text: str) -> dict[int, str]:
    """解析 '1=SMT, 2=DIP, 3=组装' 为 {阶段: 名称}。"""
    pairs = re.findall(r"(\d+)\s*=\s*([^,，]+)", str(text))
    return {int(j): name.strip() for j, name in pairs}


def _parse_weights(text: str) -> tuple[float, float, float]:
    """解析 '(0.50, 0.30, 0.20)' 为三元组。"""
    nums = [float(x) for x in re.findall(r"[\d.]+", str(text))]
    return (nums[0], nums[1], nums[2])


# ── V-toy-1 金标读入器 ────────────────────────────────────────

SHEET_PARAMS = "1_系统与算法参数"
SHEET_PLANNING = "2_计划层数据"
SHEET_SCHEDULING = "3_调度层数据"
SHEET_ORDERS = "4_订单流"


def load_vtoy1(path: str | Path) -> ProblemData:
    """读入 V-toy-1 金标工作簿，构造并校验 ProblemData。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws1 = wb[SHEET_PARAMS]
    ws2 = wb[SHEET_PLANNING]
    ws3 = wb[SHEET_SCHEDULING]
    ws4 = wb[SHEET_ORDERS]

    # —— '1_系统与算法参数'：规模与滚动机制（C3..C14）——
    products = [x.strip() for x in str(ws1["C3"].value).split(",")]
    stage_names = _parse_stage_names(ws1["C4"].value)
    stages = sorted(stage_names)
    machines = _parse_machines(ws1["C5"].value)
    speed_gears = _parse_int_set(ws1["C6"].value)
    speed_factor = _parse_gear_map(ws1["C7"].value)
    skill_levels = _parse_int_set(ws1["C8"].value)
    t_max = int(ws1["C9"].value)
    periods = list(range(1, t_max + 1))
    horizon_window = int(ws1["C10"].value)
    horizon_freeze = int(ws1["C11"].value)
    t_avail = float(ws1["C12"].value)
    ot_max = float(ws1["C13"].value)
    transport_time = float(ws1["C14"].value)

    # —— '1_系统与算法参数'：反馈与算法参数（C18..C28）——
    feedback_weights = _parse_weights(ws1["C18"].value)
    gamma = float(ws1["C19"].value)
    delta_max = float(ws1["C20"].value)
    n_restart = int(ws1["C21"].value)
    kappa_c = float(ws1["C22"].value)
    kappa_l = float(ws1["C23"].value)
    kappa_e = float(ws1["C24"].value)
    kappa_s = float(ws1["C25"].value)
    rho_min = float(ws1["C26"].value)
    epsilon_converge = float(ws1["C27"].value)
    k_max = int(ws1["C28"].value)

    # —— '2_计划层数据'：成本与批量（行 3..5）、名义产能（行 17..19）——
    cost_prod: dict[str, float] = {}
    cost_inv: dict[str, float] = {}
    cost_back: dict[str, float] = {}
    cost_setup: dict[str, float] = {}
    lot_size_max: dict[str, int] = {}
    for row in range(3, 3 + len(products)):
        p = str(ws2[f"A{row}"].value)
        cost_prod[p] = float(ws2[f"B{row}"].value)
        cost_inv[p] = float(ws2[f"C{row}"].value)
        cost_back[p] = float(ws2[f"D{row}"].value)
        cost_setup[p] = float(ws2[f"E{row}"].value)
        lot_size_max[p] = int(ws2[f"F{row}"].value)
    capacity: dict[int, dict[int, float]] = {}
    for offset, j in enumerate(stages):
        cap = float(ws2[f"C{17 + offset}"].value)
        capacity[j] = {t: cap for t in periods}  # 全部周期 A_{j,f,t}=1

    # —— '3_调度层数据'：pt^unit（行 3..5）、SDST（行 11..14）、能耗、工人 ——
    proc_time: dict[str, dict[int, float]] = {}
    for row in range(3, 3 + len(products)):
        p = str(ws3[f"A{row}"].value)
        proc_time[p] = {
            j: float(ws3.cell(row=row, column=2 + offset).value)
            for offset, j in enumerate(stages)
        }
    initial_setup: dict[str, float] = {
        p: float(ws3.cell(row=11, column=2 + i).value) for i, p in enumerate(products)
    }
    sdst: dict[str, dict[str, float]] = {}
    for i, prev in enumerate(products):
        sdst[prev] = {
            nxt: float(ws3.cell(row=12 + i, column=2 + jj).value)
            for jj, nxt in enumerate(products)
        }
    energy = EnergyParams(
        power_proc=_parse_gear_map(ws3["B20"].value),
        power_setup=float(ws3["B21"].value),
        power_idle=float(ws3["B22"].value),
        energy_transport=float(ws3["B23"].value),
        power_aux=float(ws3["B24"].value),
    )
    worker_wage: dict[int, float] = {}
    worker_available: dict[int, dict[int, int]] = {}
    for i, level in enumerate(skill_levels):
        row = 28 + i
        worker_wage[level] = float(ws3[f"B{row}"].value)
        worker_available[level] = {
            tau: int(ws3.cell(row=row, column=3 + k).value)
            for k, tau in enumerate(periods)
        }

    # —— '4_订单流'：行 3 起，直到 A 列为空 ——
    orders: list[Order] = []
    row = 3
    while True:
        oid = ws4[f"A{row}"].value
        if oid is None or not str(oid).startswith("o"):
            break
        orders.append(
            Order(
                order_id=str(oid),
                product=str(ws4[f"B{row}"].value),
                quantity=int(ws4[f"C{row}"].value),
                due_period=int(ws4[f"D{row}"].value),
                arrival_period=int(ws4[f"E{row}"].value),
            )
        )
        row += 1

    problem = ProblemData(
        products=products,
        periods=periods,
        stages=stages,
        stage_names=stage_names,
        machines=machines,
        speed_gears=speed_gears,
        skill_levels=skill_levels,
        cost_prod=cost_prod,
        cost_inv=cost_inv,
        cost_back=cost_back,
        cost_setup=cost_setup,
        lot_size_max=lot_size_max,
        capacity=capacity,
        init_inventory={p: 0.0 for p in products},  # V-toy 从零启动
        t_avail=t_avail,
        ot_max=ot_max,
        orders=orders,
        proc_time=proc_time,
        speed_factor=speed_factor,
        sdst=sdst,
        initial_setup=initial_setup,
        worker_wage=worker_wage,
        worker_available=worker_available,
        transport_time=transport_time,
        energy=energy,
        horizon_window=horizon_window,
        horizon_freeze=horizon_freeze,
        k_max=k_max,
        epsilon_converge=epsilon_converge,
        feedback_weights=feedback_weights,
        gamma=gamma,
        delta_max=delta_max,
        n_restart=n_restart,
        kappa_c=kappa_c,
        kappa_l=kappa_l,
        kappa_e=kappa_e,
        kappa_s=kappa_s,
        rho_min=rho_min,
    )
    validate_problem_data(problem)
    return problem


# ── Y企业模板结构读入骨架（字段映射留待 Stage 6）─────────────────


@dataclass(frozen=True)
class TemplateTable:
    """模板中的一张数据表：'表x-y ...' 标题 + 其下一行表头。"""

    title: str
    headers: list[str]


def load_y_template_structure(path: str | Path) -> dict[str, list[TemplateTable]]:
    """扫描 Y企业模板各 sheet，返回 {sheet 名: [表结构]}（仅结构，不读数据）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    structure: dict[str, list[TemplateTable]] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        tables: list[TemplateTable] = []
        for i, row in enumerate(rows):
            first = row[0] if row else None
            if isinstance(first, str) and re.match(r"表\d+-\d+", first):
                headers = (
                    [str(c) for c in rows[i + 1] if c is not None]
                    if i + 1 < len(rows)
                    else []
                )
                tables.append(TemplateTable(title=first, headers=headers))
        structure[ws.title] = tables
    return structure
